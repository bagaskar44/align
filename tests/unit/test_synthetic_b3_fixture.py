from __future__ import annotations

import importlib.util
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
FIXTURE = ROOT / "fixtures" / "b3" / "complex-submission.xlsx"
GENERATOR = ROOT / "scripts" / "generate-b3-fixtures.py"


def _formula_cells(path: Path) -> dict[str, dict[str, str]]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    return {
        sheet.title: {
            cell.coordinate: cell.value
            for row in sheet.iter_rows()
            for cell in row
            if cell.data_type == "f"
        }
        for sheet in workbook.worksheets
        if any(cell.data_type == "f" for row in sheet.iter_rows() for cell in row)
    }


def test_fixture_generator_matches_local_workbook_sheets_and_unverified_formulas(tmp_path: Path) -> None:
    output = tmp_path
    spec = importlib.util.spec_from_file_location("b3_fixture_generator", GENERATOR)
    assert spec is not None and spec.loader is not None
    generator = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(generator)
    generator.generate_xlsx(output)
    generated = output / "complex-submission.xlsx"

    original_book = load_workbook(FIXTURE, data_only=False)
    generated_book = load_workbook(generated, data_only=False)
    assert generated_book.sheetnames == original_book.sheetnames
    assert {sheet.title: sheet.sheet_state for sheet in generated_book.worksheets} == {
        sheet.title: sheet.sheet_state for sheet in original_book.worksheets
    }
    assert original_book["Rekap Nilai Semester Akademik"]["A1"].value == "Data untuk tab dengan nama panjang"
    assert generated_book["Rekap Nilai Semester Akademik"]["A1"].value == "Data untuk tab dengan nama panjang"
    assert _formula_cells(generated) == _formula_cells(FIXTURE)
    assert len(_formula_cells(generated)["Ringkasan"]) == 2
    assert original_book["Catatan internal"].sheet_state == "hidden"

    original_cached = load_workbook(FIXTURE, data_only=True)
    generated_cached = load_workbook(generated, data_only=True)
    for coordinate in ("D3", "D4"):
        assert original_cached["Ringkasan"][coordinate].value is None
        assert generated_cached["Ringkasan"][coordinate].value is None
