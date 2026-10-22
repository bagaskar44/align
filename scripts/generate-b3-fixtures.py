"""Generate deterministic-content synthetic OOXML fixtures with locked parser libraries."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference


def generate_docx(output: Path) -> None:
    from docx import Document

    document = Document()
    document.sections[0].header.paragraphs[0].text = "Header sintetis: B3 parsing source"
    document.add_paragraph("Ringkasan sintetis: seluruh angka merupakan fixture uji.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Bagian"
    table.cell(0, 1).text = "Isi"
    row = table.add_row().cells
    row[0].text = "Konteks"
    row[1].text = "Narasi sintetis"
    nested = row[1].add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Fakta"
    nested.cell(0, 1).text = "Sumber"
    nested.add_row().cells[0].text = "Contoh"
    document.add_paragraph("Rujuk tabel di atas untuk struktur yang dapat ditinjau.")
    document.sections[0].footer.paragraphs[0].text = "Footer sintetis"
    document.save(output / "complex-submission.docx")


def generate_xlsx(output: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Ringkasan"
    sheet.merge_cells("A1:D1")
    sheet["A1"] = "Laporan sintetis — bukan data mahasiswa"
    sheet.append(["Periode", "Unit A", "Unit B", "Total"])
    sheet.append(["2026-Q1", 12, 8, "=SUM(B3:C3)"])
    sheet.append(["2026-Q2", 15, 11, "=SUM(B4:C4)"])
    sheet["B3"].number_format = '#,##0 "unit"'
    sheet.column_dimensions["A"].width = 34
    sheet.sheet_view.showGridLines = False
    hidden = workbook.create_sheet("Catatan internal")
    hidden.sheet_state = "hidden"
    hidden.append(["Kode", "Definisi"])
    hidden.append(["SYN-01", "Nilai sintetis untuk uji inventory sheet tersembunyi"])
    chart = BarChart()
    chart.title = "Perbandingan sintetis"
    chart.add_data(Reference(sheet, min_col=2, max_col=3, min_row=2, max_row=4), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=3, max_row=4))
    sheet.add_chart(chart, "F2")
    semester = workbook.create_sheet("Rekap Nilai Semester Akademik")
    semester["A1"] = "Data untuk tab dengan nama panjang"
    workbook.save(output / "complex-submission.xlsx")


def main() -> None:
    output = Path(sys.argv[1] if len(sys.argv) > 1 else "fixtures/b3")
    output.mkdir(parents=True, exist_ok=True)
    generate_docx(output)
    generate_xlsx(output)


if __name__ == "__main__":
    main()
